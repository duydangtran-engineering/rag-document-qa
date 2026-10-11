"""Read documents from disk and return text blocks that remember where they came from.

The page number is the whole point: without it we cannot show a real citation later.

What "clean ingestion" means here (each rule exists because of a real defect found in
the first 450 chunks):

  * running headers / footers and page numbers are removed  -> otherwise the line
    "Tổng quan và phân tích chính sách bảo hiểm xã hội" was embedded in ~80 chunks and
    made every chunk of the report look a little alike.
  * table-of-contents / list-of-figures pages are dropped    -> they contain titles of
    sections, so they win similarity searches without containing any answer.
  * tables are read as tables (header row kept, one row per line) instead of a soup
    of cells -> "Bộ Tài chính" was an answer with no column name next to it.
  * heading detection is strict: a heading must look like a heading, not merely start
    with the word "Chương" or an upper-case word.
"""

from __future__ import annotations

import math
import re
from collections import Counter
from dataclasses import dataclass
from pathlib import Path

import pymupdf  # the `fitz` alias is deprecated
from docx import Document as DocxDocument


@dataclass
class PageBlock:
    """One unit of raw text plus its location in the source file."""

    text: str
    page: int  # 1-based. For DOCX/TXT we use a pseudo page number.
    source: str  # file name


# --------------------------------------------------------------------------- #
# Headings
# --------------------------------------------------------------------------- #

# A PDF gives us hard line wraps in the middle of sentences and NO blank lines
# between paragraphs. Splitting on "\n\n" therefore returns the whole page as one
# block, which defeats paragraph-aware chunking and destroys heading detection.
# So we rebuild paragraph boundaries from the numbering markers that Vietnamese
# legal / policy documents use at the start of every clause.
LINE_START_MARKER = re.compile(
    r"^[“\"']?\s*(?:"
    r"(?:Điều|ĐIỀU|Chương|CHƯƠNG|Mục|MỤC|Phần|PHẦN|Phụ\s+lục|PHỤ\s+LỤC)\s+[0-9IVXLC]+"
    r"|\d+[a-zđ]?\s*[.)]"      # 1.   4)   3a.   (khoản)
    r"|\d+\.\d+"               # 3.2        (mục trong phụ lục)
    r"|[a-zđ]\s*\)"            # a)   đ)    (điểm)
    r")"
)

MAX_HEADING_LINE = 120

# "Điều 4. Điều khoản chuyển tiếp", "PHẦN II: GIỚI THIỆU", "Chương IV", "PHỤ LỤC 3 PHƯƠNG PHÁP"
# The number must be followed by a delimiter (. : – -) or the end of the line.
# That rejects a wrapped line of body text such as "Chương IV, Luật Bảo hiểm xã hội."
# or "Điều 18 như sau:", which the old pattern happily accepted as headings.
_STRUCT = re.compile(
    r"^[“\"']?\s*(Điều|ĐIỀU|Chương|CHƯƠNG|Mục|MỤC|Phần|PHẦN|PART|Part|Phụ\s*lục|PHỤ\s*LỤC)"
    r"\s*([0-9]{1,3}|[IVXLC]{1,6})\b\s*(?P<rest>.*)$"
)
_ROMAN = re.compile(r"^([IVX]{1,4})\s*[.)]\s+(?P<rest>\S.*)$")
# 1.1  1.1.2.  — each part at most 2 digits and the first part is not 0, so a
# measurement such as "0.1964 mean_S" or a year "2.000" is never a section number.
_DECIMAL = re.compile(r"^([1-9][0-9]?(?:\.[0-9]{1,2})+)\.?\s+(?P<rest>\S.*)$")
_NUMBERED = re.compile(r"^([0-9]{1,2})\s*[.)]\s+(?P<rest>\S.*)$")

_STRUCT_LEVEL = {"phần": 1, "part": 1, "chương": 1, "phụ lục": 1, "mục": 2, "điều": 3}

# Levels, used to build the "where am I in the document" path:
#   1 Phần / Chương / Phụ lục     2 Mục, I., unnumbered CAPITALS     3 Điều, 1. CAPITALS
#   4+ 1.1 / 1.1.2     (the "(được sửa đổi)" label of an amending clause is level 3)
AMEND_LEVEL = 3  # same level as Điều: an amended article replaces the one before it


def _letters(text: str) -> list[str]:
    return [c for c in text if c.isalpha()]


def _upper_ratio(text: str) -> float:
    letters = _letters(text)
    return sum(1 for c in letters if c.isupper()) / len(letters) if letters else 0.0


def _looks_like_caps_title(text: str) -> bool:
    """ALL-CAPS section title: several words, no code / formula / table-header look."""
    text = text.strip()
    if len(_letters(text)) < 8 or _upper_ratio(text) < 0.85:
        return False
    if len(text.split()) < 2:  # "MANHATTAN" is a word, not a section.
        return False
    if any(ch in text for ch in "_=/()[]{}<>|*%"):  # BRIGHT_THRESHOLD, (DY/DX), ...
        return False
    if re.search(r"\s\d{1,3}$", text) or text.endswith((",", ";")):  # TOC row / mid-sentence
        return False
    digits = sum(c.isdigit() for c in text)
    return digits <= 0.25 * len(_letters(text))


def parse_heading(line: str) -> tuple[int, str, str] | None:
    """(level, title, kind) when the line is a standalone heading, else None.

    kind is "struct" for lines with a structural prefix (Điều 5., II., 1.2.) and
    "caps" for an unnumbered ALL-CAPS title. A "caps" line directly under a heading
    is usually its wrapped second line ("Chương II" / "QUY HOẠCH ĐÔ THỊ").

    The length cap matters as much as the pattern: a body sentence such as
    "theo Điều 5, Chương I, Luật BHXH năm 2014: Mức đóng..." also contains
    the word "Điều" but is long and has a comma right after the number.
    """
    line = line.strip()
    if not line or len(line) > MAX_HEADING_LINE:
        return None
    title = line.strip("“\"'").rstrip(".:").strip()
    title = re.sub(r"^([^\W\d_]+)(\d)", r"\1 \2", title)  # "PHẦN2" -> "PHẦN 2"

    # Dotted leaders belong to a table of contents, never to a heading.
    if re.search(r"\.{4,}|…{2,}", line):
        return None

    m = _STRUCT.match(line)
    if m:
        keyword = re.sub(r"\s+", " ", m.group(1).lower())
        rest = m.group("rest")
        after_number = line[m.end(2):]
        delimited = rest == "" or bool(re.match(r"\s*[.:–—-]", after_number))
        # "Chương II QUY HOẠCH ĐÔ THỊ" (two wrapped lines merged) or "PHỤ LỤC 3 PHƯƠNG PHÁP"
        shouted = rest != "" and rest[:1].isupper() and _upper_ratio(rest) >= 0.85
        if delimited or shouted:
            return _STRUCT_LEVEL[keyword], title, "struct"
        return None

    m = _ROMAN.match(line)
    if m and m.group("rest")[0].isupper():
        return 2, title, "struct"

    m = _DECIMAL.match(line)
    if m and m.group("rest")[0].isupper():
        return 3 + m.group(1).count("."), title, "struct"

    m = _NUMBERED.match(line)
    if m and _looks_like_caps_title(m.group("rest")):
        return 3, title, "struct"

    if _looks_like_caps_title(line):
        return 2, title, "caps"
    return None


def is_heading_line(line: str) -> bool:
    return parse_heading(line) is not None


# --------------------------------------------------------------------------- #
# Text reflow
# --------------------------------------------------------------------------- #

_LEADER = re.compile(r"\.{5,}|…{2,}|(?:\. ){5,}")


def _reflow(text: str) -> str:
    """Join wrapped lines back into paragraphs, one paragraph per clause."""
    paragraphs: list[str] = []
    current: list[str] = []
    previous_was_heading = False

    def flush() -> None:
        if current:
            paragraphs.append(" ".join(current))
            current.clear()

    for raw_line in text.split("\n"):
        line = raw_line.strip()
        if not line:
            flush()  # A real blank line is always a paragraph break.
            previous_was_heading = False
            continue

        heading = parse_heading(line)
        if heading:
            flush()
            # A long heading wraps onto a second line. Only an unnumbered CAPITALS
            # line is merged into the heading above it: two consecutive headings
            # such as "PHẦN II: ..." and "I. KHUNG PHÁP LÝ" are different levels
            # and must stay separate.
            if (previous_was_heading and heading[2] == "caps"
                    and len(paragraphs[-1]) + len(line) + 1 <= MAX_HEADING_LINE):
                paragraphs[-1] = f"{paragraphs[-1]} {line}"
            else:
                paragraphs.append(line)  # Headings stand alone, never absorbed.
            previous_was_heading = True
            continue

        if current and LINE_START_MARKER.match(line):
            flush()  # A new numbered clause starts here.
        current.append(line)
        previous_was_heading = False

    flush()
    return "\n\n".join(p for p in paragraphs if p)


def _clean(text: str) -> str:
    """Normalise whitespace, then rebuild paragraph structure."""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    return _reflow(text).strip()


# --------------------------------------------------------------------------- #
# PDF: tables
# --------------------------------------------------------------------------- #

MAX_TABLE_COLS = 12
MIN_TABLE_FILL = 0.55  # share of non-empty cells; a chart mistaken for a table is mostly empty


def _cell(value) -> str:
    return re.sub(r"\s+", " ", (value or "").replace("\n", " ")).strip()


def table_to_markdown(rows: list[list]) -> str | None:
    """Render a PyMuPDF table as a Markdown table, or None if it is not a real table.

    * empty rows / columns are dropped (merged cells leave plenty of them)
    * a row whose first cell is empty continues the row above it (a cell that wraps
      over several lines is returned by PyMuPDF as several rows)
    * the header row stays the first line, and the chunker repeats it in every
      chunk the table is split into, so a cell never loses its column name
    """
    grid = [[_cell(c) for c in row] for row in rows if row]
    grid = [r for r in grid if any(r)]
    if len(grid) < 2:
        return None
    width = max(len(r) for r in grid)
    grid = [r + [""] * (width - len(r)) for r in grid]
    keep = [j for j in range(width) if any(r[j] for r in grid)]
    grid = [[r[j] for j in keep] for r in grid]
    width = len(keep)
    if width < 2 or width > MAX_TABLE_COLS:
        return None

    merged: list[list[str]] = []
    for i, row in enumerate(grid):
        if i > 0 and merged and not row[0]:
            for j, value in enumerate(row):
                if value:
                    merged[-1][j] = f"{merged[-1][j]} {value}".strip()
        else:
            merged.append(list(row))
    if len(merged) < 2:
        return None

    filled = sum(1 for r in merged for c in r if c)
    if filled / (len(merged) * width) < MIN_TABLE_FILL:
        return None
    # A table of bare numbers with no words is a chart's data labels.
    if sum(1 for r in merged for c in r if any(ch.isalpha() for ch in c)) < 2:
        return None

    def line(cells: list[str]) -> str:
        return "| " + " | ".join(c.replace("|", "/") for c in cells) + " |"

    out = [line(merged[0]), "| " + " | ".join("---" for _ in range(width)) + " |"]
    out.extend(line(r) for r in merged[1:])
    return "\n".join(out)


def _good_tables(page) -> list[tuple[tuple, str]]:
    try:
        found = page.find_tables().tables
    except Exception:  # old PyMuPDF, or an odd page: fall back to plain text
        return []
    tables = []
    for table in found:
        md = table_to_markdown(table.extract())
        if md:
            tables.append((tuple(table.bbox), md))
    return tables


# --------------------------------------------------------------------------- #
# PDF: one page -> items, then document-level cleaning
# --------------------------------------------------------------------------- #

def _page_items(page) -> list[list]:
    """The page as an ordered list of ["text", str] / ["table", markdown].

    Blocks are read in the PDF's own reading order (NOT sorted by position: this
    report is two-column and sorting by y would interleave the columns). A table
    takes the place of its first block; the other blocks inside it are skipped.
    """
    tables = _good_tables(page)
    items: list[list] = []
    used: set[int] = set()
    for block in page.get_text("blocks"):
        if block[6] != 0:  # image block
            continue
        cx, cy = (block[0] + block[2]) / 2, (block[1] + block[3]) / 2
        inside = next(
            (i for i, (bb, _) in enumerate(tables)
             if bb[0] - 2 <= cx <= bb[2] + 2 and bb[1] - 2 <= cy <= bb[3] + 2),
            None,
        )
        if inside is None:
            items.append(["text", block[4]])
        elif inside not in used:
            used.add(inside)
            items.append(["table", tables[inside][1]])
    items.extend(["table", md] for i, (_, md) in enumerate(tables) if i not in used)
    return items


_PAGE_NUMBER = re.compile(r"^\d{1,3}$")
EDGE_LINES = 3  # how many lines at the top / bottom of a page count as "edge"


def _running_key(line: str) -> str:
    """A line with its page number removed, so '4 Tên báo cáo' and 'Tên báo cáo 5' match."""
    line = re.sub(r"^\d{1,3}\s+|\s+\d{1,3}$", "", line.strip())
    return re.sub(r"\d+", "#", re.sub(r"\s+", " ", line))


def _all_lines(items: list[list]) -> list[tuple[int, int, str]]:
    """(item index, line index, text) of every non-empty text line on the page."""
    return [
        (i, j, ln.strip())
        for i, (kind, content) in enumerate(items) if kind == "text"
        for j, ln in enumerate(content.split("\n")) if ln.strip()
    ]


def _edge_positions(items: list[list]) -> list[tuple[int, int, str]]:
    """The first/last lines of the page's text."""
    lines = _all_lines(items)
    if len(lines) <= 2 * EDGE_LINES:
        return lines
    return lines[:EDGE_LINES] + lines[-EDGE_LINES:]


def strip_running_lines(pages: list[list[list]]) -> list[str]:
    """Remove page numbers and repeated header/footer lines. Returns what was removed."""
    if not pages:
        return []
    counts: Counter = Counter()
    for items in pages:
        for key in {_running_key(t) for _, _, t in _edge_positions(items)}:
            counts[key] += 1
    threshold = max(3, math.ceil(0.3 * len(pages)))
    repeated = {k for k, n in counts.items()
                if n >= threshold and len(pages) >= 4 and k and k != "#"}

    for items in pages:
        lines = _all_lines(items)
        drop = {(i, j) for i, j, t in _edge_positions(items) if _PAGE_NUMBER.match(t)}
        for position, (i, j, t) in enumerate(lines):
            if _running_key(t) not in repeated:
                continue
            # The header is not always among the first lines in reading order (a
            # two-column page, a title block): remove it wherever it is, together
            # with the page number printed right beside it.
            drop.add((i, j))
            for k in (position - 1, position + 1):
                if 0 <= k < len(lines) and _PAGE_NUMBER.match(lines[k][2]):
                    drop.add(lines[k][:2])
        if not drop:
            continue
        for i, item in enumerate(items):
            if item[0] != "text":
                continue
            item[1] = "\n".join(
                ln for j, ln in enumerate(item[1].split("\n")) if (i, j) not in drop)
    return sorted(repeated)


_TOC_TITLE = re.compile(
    r"^(MỤC\s+LỤC|DANH\s+MỤC.*|TABLE\s+OF\s+CONTENTS|CONTENTS)$", re.IGNORECASE)


def _is_index_line(line: str) -> bool:
    line = line.strip()
    return bool(
        re.fullmatch(r"\d{1,3}", line)                       # bare page number
        or _LEADER.search(line)                              # "Tiêu đề ........ 12"
        or re.search(r"[^\d\s]\s*\.{2,}\s*\d{1,3}$", line)
    )


_CAPTION = re.compile(r"^(Hình|Bảng|Biểu đồ|Figure|Table)\s*\d+\s*[.:]", re.IGNORECASE)
_BARE_CAPTION = re.compile(r"^(Hình|Bảng|Biểu đồ|Figure|Table)\s*\d+\s*[.:]?\s*$", re.IGNORECASE)
MAX_PROSE_LINES = 2
MIN_LIST_CAPTIONS = 5  # this many "Hình N." lines on one page = a list of figures


def _is_prose_line(line: str) -> bool:
    """A line of running text: several words, mostly lower case, not a caption."""
    if len(line.split()) < 6 or _CAPTION.match(line.strip()) or _LEADER.search(line):
        return False
    return _upper_ratio(line) < 0.3


def is_index_page(items: list[list]) -> bool:
    """A table of contents / list of figures: mostly titles and page numbers.

    Such a page is full of section names, so it ranks high for questions that
    mention a topic, yet it never contains the answer. Only pages where a large
    share of the lines are bare numbers or dotted leaders are dropped, so a normal
    page that merely mentions a page number is safe.
    """
    lines = [ln.strip() for kind, c in items if kind == "text" for ln in c.split("\n") if ln.strip()]
    if len(lines) < 8:
        return False
    # A chart page is also full of bare numbers (axis labels) but it has sentences
    # around the chart; a table of contents has none.
    # A list of figures / tables: a column of "Hình N." / "Bảng N." entries.
    if sum(bool(_CAPTION.match(ln)) for ln in lines) >= MIN_LIST_CAPTIONS:
        return True
    prose = 0
    for previous, line in zip([""] + lines, lines):
        # "Hình 7." and its title are often extracted as two separate lines.
        if _BARE_CAPTION.match(previous):
            continue
        prose += _is_prose_line(line)
    if prose > MAX_PROSE_LINES:
        return False
    share = sum(_is_index_line(ln) for ln in lines) / len(lines)
    titled = any(_TOC_TITLE.match(ln) for ln in lines[:4])
    return share >= (0.18 if titled else 0.3)


def _render_page(items: list[list]) -> str:
    """Reflow the text runs, keep tables intact, join everything with blank lines."""
    parts: list[str] = []
    run: list[str] = []

    def flush_run() -> None:
        if run:
            text = "\n".join(ln for ln in "\n".join(run).split("\n") if not _LEADER.search(ln))
            cleaned = _clean(text)
            if cleaned:
                parts.append(cleaned)
            run.clear()

    for kind, content in items:
        if kind == "text":
            run.append(content)
        else:
            flush_run()
            parts.append(content)
    flush_run()
    return "\n\n".join(parts)


def load_pdf_with_report(path: Path) -> tuple[list[PageBlock], dict]:
    with pymupdf.open(path) as doc:
        pages = [_page_items(page) for page in doc]

    running = strip_running_lines(pages)
    blocks: list[PageBlock] = []
    index_pages: list[int] = []
    for number, items in enumerate(pages, start=1):
        if is_index_page(items):
            index_pages.append(number)
            continue
        text = _render_page(items)
        if text:
            blocks.append(PageBlock(text=text, page=number, source=path.name))
    report = {
        "pages": len(pages),
        "kept_pages": len(blocks),
        "index_pages": index_pages,
        "running_lines": running,
        "tables": sum(1 for items in pages for kind, _ in items if kind == "table"),
    }
    return blocks, report


def load_pdf(path: Path) -> list[PageBlock]:
    return load_pdf_with_report(path)[0]


# --------------------------------------------------------------------------- #
# DOCX / TXT / dispatch
# --------------------------------------------------------------------------- #

def load_docx(path: Path) -> list[PageBlock]:
    """DOCX has no real pages. We group paragraphs into pseudo-pages of ~3000 chars
    so citations still point at a narrow, findable region of the file."""
    doc = DocxDocument(path)

    parts: list[str] = [p.text for p in doc.paragraphs if p.text.strip()]

    # Tables often hold the actual policy numbers, so do not skip them.
    for table in doc.tables:
        rows = [[_cell(c.text) for c in row.cells] for row in table.rows]
        markdown = table_to_markdown(rows)
        if markdown:
            parts.append(markdown)
            continue
        for row in rows:
            cells = [c for c in row if c]
            if cells:
                parts.append(" | ".join(cells))

    blocks: list[PageBlock] = []
    buffer: list[str] = []
    size = 0
    pseudo_page = 1

    def make(texts: list[str], page: int) -> PageBlock:
        # Markdown tables must survive _clean (which reflows lines), so keep them apart.
        out: list[str] = []
        run: list[str] = []
        for t in texts:
            if t.startswith("|"):
                if run:
                    out.append(_clean("\n\n".join(run)))
                    run = []
                out.append(t)
            else:
                run.append(t)
        if run:
            out.append(_clean("\n\n".join(run)))
        return PageBlock(text="\n\n".join(o for o in out if o), page=page, source=path.name)

    for part in parts:
        buffer.append(part)
        size += len(part)
        if size >= 3000:
            blocks.append(make(buffer, pseudo_page))
            buffer, size = [], 0
            pseudo_page += 1

    if buffer:
        blocks.append(make(buffer, pseudo_page))
    return [b for b in blocks if b.text]


def load_text(path: Path) -> list[PageBlock]:
    text = _clean(path.read_text(encoding="utf-8", errors="ignore"))
    return [PageBlock(text=text, page=1, source=path.name)] if text else []


def load_document(path: Path) -> list[PageBlock]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return load_pdf(path)
    if suffix == ".docx":
        return load_docx(path)
    if suffix in {".txt", ".md"}:
        return load_text(path)
    raise ValueError(f"Unsupported file type: {path.name}")


def load_directory(directory: Path, extensions: set[str]) -> list[PageBlock]:
    blocks: list[PageBlock] = []
    for path in sorted(directory.rglob("*")):
        if path.is_file() and path.suffix.lower() in extensions:
            blocks.extend(load_document(path))
    return blocks
